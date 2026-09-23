from __future__ import annotations

from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import time as _time
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .models import Event, Operation, ParsedRow, ParsedWorkbook, RowIssue


HEADERS = (
    "№ задания",
    "Стикер",
    "КИЗ",
    "Номер чека",
    "Стоимость",
    "Валюта",
    "Номер фискального накопителя",
    "Дата",
    "Тип операции",
    "Признак продажи юрлицу",
)

PARSER_VERSION = 2

# Explicit project assumption, not inferred from the Windows local timezone.
WB_TIMEZONE = timezone(timedelta(hours=3))
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_UNPACKED_BYTES = 250 * 1024 * 1024
MAX_ROWS = 200_000
MAX_ZIP_ENTRIES = 4096
MAX_ZIP_ENTRY_BYTES = 100 * 1024 * 1024
MAX_HEADER_COLUMNS = 256
MAX_PARSE_SECONDS = 30.0


class WorkbookError(ValueError):
    """Whole-file problem: unsupported/corrupt workbook or wrong schema."""


class RowValidationError(ValueError):
    """A rejected row; other rows can still be imported."""


def _text(value: str) -> str:
    # Do NOT use bare str.strip(): it can remove ASCII GS (\x1d).
    return value.strip(" \t\r\n")


def _blank(value: Any) -> bool:
    return value is None or isinstance(value, str) and not _text(value)


def _header(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(value.replace("\xa0", " ").split())


def _identifier(
    value: Any, field: str, *, required: bool = True, text_only: bool = False,
) -> str | None:
    if _blank(value):
        if required:
            raise RowValidationError(f"{field}: значение обязательно")
        return None
    if isinstance(value, str):
        return _text(value)
    if text_only:
        raise RowValidationError(
            f"{field}: нужен текстовый Excel-формат; числовой код небезопасен"
        )
    # Excel stores at most 15 reliable significant decimal digits.
    if type(value) is int and 0 <= value < 10**15:
        return str(value)
    if type(value) is float and value.is_integer() and 0 <= value < 10**15:
        return str(int(value))
    raise RowValidationError(
        f"{field}: нужен текст либо целое число менее 16 цифр; "
        "проверьте потерю точности Excel"
    )


def _amount(value: Any) -> Decimal:
    if _blank(value) or isinstance(value, bool):
        raise RowValidationError("Стоимость: требуется число")
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        result = Decimal(text)
        if not result.is_finite() or abs(result) > Decimal("999999999999999"):
            raise RowValidationError("Стоимость: некорректное значение")
        rounded = result.quantize(Decimal("0.01"))
        if result != rounded:
            raise RowValidationError("Стоимость: более двух знаков после запятой")
        return abs(rounded) if rounded == 0 else rounded
    except InvalidOperation as exc:
        raise RowValidationError("Стоимость: не удалось прочитать число") from exc


def _currency(value: Any) -> str:
    if not isinstance(value, str):
        raise RowValidationError("Валюта: требуется текстовый код")
    normalized = _text(value).upper()
    if normalized in {"RUR", "РУБ", "РУБ.", "₽"}:
        normalized = "RUB"
    if not re.fullmatch(r"[A-Z]{3}", normalized):
        raise RowValidationError("Валюта: ожидается трёхбуквенный код, например RUB")
    return normalized


def _date(value: Any) -> datetime | None:
    # An absent WB date is not an error and must never acquire an invented time.
    if _blank(value):
        return None

    parsed: datetime | None = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        # Retained legacy support for an explicitly supplied calendar date.
        # This is NOT used for absent dates; midnight is not evidence of order.
        parsed = datetime.combine(value, time())
    elif isinstance(value, str):
        text = _text(value)
        try:
            parsed = datetime.fromisoformat(text)
        except ValueError:
            for fmt in (
                "%H:%M:%S %d.%m.%Y",
                "%d.%m.%Y %H:%M:%S",
                "%d.%m.%Y %H:%M",
                "%d.%m.%Y",
            ):
                try:
                    parsed = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue

    if parsed is None:
        raise RowValidationError(
            "Дата: ожидается пустое значение, Excel-дата, ISO-дата, "
            "ЧЧ:ММ:СС ДД.ММ.ГГГГ или ДД.ММ.ГГГГ [ЧЧ:ММ:СС]"
        )
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        parsed = parsed.replace(tzinfo=WB_TIMEZONE)
    return parsed.astimezone(timezone.utc)


def _legal(value: Any) -> bool | None:
    if _blank(value):
        return None
    if type(value) is bool:
        return value
    if type(value) in (int, float) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        text = _text(value).casefold()
        if text in {"да", "true", "1"}:
            return True
        if text in {"нет", "false", "0"}:
            return False
    raise RowValidationError("Признак продажи юрлицу: ожидается да/нет или 1/0")


def _parse_row(values: dict[str, Any]) -> Event:
    operation_text = values["Тип операции"]
    if not isinstance(operation_text, str):
        raise RowValidationError("Тип операции: ожидается Продажа или Возврат")
    mapping = {"продажа": Operation.SALE, "возврат": Operation.RETURN}
    operation = mapping.get(_text(operation_text).casefold())
    if operation is None:
        raise RowValidationError(
            f"Тип операции: неизвестное значение {operation_text!r}"
        )
    kiz = _identifier(values["КИЗ"], "КИЗ", text_only=True)
    task = _identifier(values["№ задания"], "№ задания")
    sticker = _identifier(values["Стикер"], "Стикер")
    assert kiz is not None and task is not None and sticker is not None
    return Event(
        kiz=kiz,
        task_number=task,
        sticker=sticker,
        operation=operation,
        occurred_at=_date(values["Дата"]),
        receipt_number=_identifier(
            values["Номер чека"], "Номер чека", required=False
        ),
        fiscal_drive_number=_identifier(
            values["Номер фискального накопителя"],
            "Номер фискального накопителя",
            required=False,
        ),
        amount=_amount(values["Стоимость"]),
        currency=_currency(values["Валюта"]),
        legal_entity_sale=_legal(values["Признак продажи юрлицу"]),
    )


def parse_excel(data: bytes) -> ParsedWorkbook:
    started = _time.monotonic()
    if len(data) > MAX_FILE_BYTES:
        raise WorkbookError("Файл превышает допустимый размер 50 MiB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise WorkbookError("Excel-архив содержит слишком много ZIP-записей")
            if sum(item.file_size for item in entries) > MAX_UNPACKED_BYTES:
                raise WorkbookError("Слишком большой распакованный Excel-архив")
            for item in entries:
                normalized = item.filename.replace("\\", "/")
                parts = [part for part in normalized.split("/") if part not in ("", ".")]
                if any(part == ".." for part in parts) or normalized.startswith("/"):
                    raise WorkbookError("Excel-архив содержит небезопасный путь")
                if item.flag_bits & 0x1:
                    raise WorkbookError("Зашифрованные ZIP-записи Excel не поддерживаются")
                if item.file_size > MAX_ZIP_ENTRY_BYTES:
                    raise WorkbookError("Excel-архив содержит чрезмерно большую запись")
                lowered = normalized.casefold()
                if lowered.startswith("xl/externallinks/") or lowered.startswith("xl/embeddings/") or lowered.endswith("vbaproject.bin"):
                    raise WorkbookError("Внешние ссылки, вложенные объекты и макросы не допускаются")
            for item in entries:
                if item.filename.casefold().endswith(".rels") and item.file_size <= 2 * 1024 * 1024:
                    rel = archive.read(item)
                    if b'TargetMode="External"' in rel or b"TargetMode='External'" in rel:
                        raise WorkbookError("Внешние связи Excel не допускаются")
    except BadZipFile as exc:
        raise WorkbookError(
            "Ожидается Excel .xlsx; старый .xls, ZIP с вложенным Excel "
            "и произвольные файлы не поддерживаются"
        ) from exc

    workbook = None
    try:
        # Formulas are deliberately NOT evaluated or silently replaced by caches.
        workbook = load_workbook(
            BytesIO(data), read_only=True, data_only=False, keep_links=False,
        )
        if "КИЗ" not in workbook.sheetnames:
            raise WorkbookError("В Excel отсутствует обязательный лист «КИЗ»")
        sheet = workbook["КИЗ"]
        # Do not trust potentially incorrect worksheet dimension metadata.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        first = next(iterator, None)
        if first is None:
            raise WorkbookError("Лист «КИЗ» пуст")
        if len(first) > MAX_HEADER_COLUMNS:
            raise WorkbookError(f"Слишком много колонок в листе «КИЗ»: максимум {MAX_HEADER_COLUMNS}")
        names = [_header(cell.value) for cell in first]
        missing = [name for name in HEADERS if name not in names]
        if missing:
            raise WorkbookError("Отсутствуют обязательные колонки: " + ", ".join(missing))
        duplicated = [name for name in HEADERS if names.count(name) > 1]
        if duplicated:
            raise WorkbookError("Повторяются колонки: " + ", ".join(duplicated))
        positions = {name: names.index(name) for name in HEADERS}
        rows: list[ParsedRow] = []
        issues: list[RowIssue] = []
        ignored_rows: list[int] = []
        for number, cells in enumerate(iterator, start=2):
            if number % 1000 == 0 and _time.monotonic() - started > MAX_PARSE_SECONDS:
                raise WorkbookError("Превышен безопасный лимит времени разбора Excel")
            if number > MAX_ROWS + 1:
                raise WorkbookError(f"Превышен лимит строк: {MAX_ROWS}")
            values = {
                name: cells[index].value if index < len(cells) else None
                for name, index in positions.items()
            }
            if all(_blank(value) for value in values.values()):
                continue
            operation_value = values.get("Тип операции")
            if isinstance(operation_value, str) and _text(operation_value) == "-":
                ignored_rows.append(number)
                continue
            try:
                for name, index in positions.items():
                    if index < len(cells) and cells[index].data_type in {"f", "e"}:
                        raise RowValidationError(
                            f"{name}: формулы и Excel-ошибки не допускаются; "
                            "сохраните исходное значение"
                        )
                rows.append(ParsedRow(number, _parse_row(values)))
            except (ValueError, TypeError, OverflowError) as exc:
                issues.append(RowIssue(number, str(exc)))
        return ParsedWorkbook(tuple(rows), tuple(issues), tuple(ignored_rows))
    except WorkbookError:
        raise
    except Exception as exc:
        # File decoding is a boundary with third-party ZIP/XML exceptions.
        raise WorkbookError(
            f"Не удалось прочитать Excel ({type(exc).__name__}): {exc}"
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
