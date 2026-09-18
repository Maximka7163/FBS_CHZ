from __future__ import annotations

import io
import json
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from wbcz.m11_reports import (
    ARBITRARY_REMOTE_URL_DOWNLOAD,
    ARTIFACT_ENCRYPTION_FORMAT,
    CALLER_CONTROLLED_FILESYSTEM_PATH,
    CALLER_CONTROLLED_STORAGE_KEY,
    CALLER_CONTROLLED_TRUE_API_HOST,
    CALLER_CONTROLLED_TRUE_API_METHOD,
    CALLER_CONTROLLED_TRUE_API_PATH,
    DISPENSER_CAPABILITIES,
    FRONTEND_FREEZE_ACTIVE,
    GENERIC_TRUE_API_REPORT_PROXY,
    KNOWN_REMOTE_AVAILABILITY,
    KNOWN_REMOTE_DOWNLOAD_STATUSES,
    KNOWN_REMOTE_TASK_STATUSES,
    LOCAL_REPORT_CATALOG,
    M11_AGENT_JOB_TYPES,
    M12_DEPENDENCY_FOR_DOWNLOAD_AUTH,
    PATH_TRAVERSAL_ALLOWED,
    PRODUCTION_TRUE_API_REPORTS_ENABLED,
    REMOTE_RESULT_DELETE_ENABLED,
    REPORT_CREATE_RATE_PER_MINUTE,
    REPORT_QUOTA_RATE_PER_MINUTE,
    REPORT_RESULTS_RATE_PER_MINUTE,
    REPORT_TASK_RATE_PER_MINUTE,
    USER_REPORT_DOWNLOAD_ENABLED,
    ArtifactIntegrityConflict,
    ArtifactLimitExceeded,
    ArtifactRole,
    BinaryArtifactIngress,
    ChunkedAeadArtifactCipher,
    DispenserCapabilityName,
    DispenserRateLimiter,
    FilesystemReportArtifactStore,
    FilteredCisFilters,
    InMemoryUploadBindingStore,
    LocalReconciliationResult,
    LocalReportType,
    ReportContractError,
    ReportFormatError,
    ReportOutputFormat,
    ReportRateLimitExceeded,
    ReportSecurityError,
    ReportSensitivity,
    ReportSnapshotDescriptor,
    SnapshotStrategy,
    StreamingReportRenderer,
    artifact_reuse_allowed,
    build_dispenser_spec,
    build_filtered_cis_create_body,
    build_reconciliation_row,
    canonical_request_fingerprint,
    disabled_remote_recipes,
    enabled_remote_recipes,
    format_timestamp_preserving_unknown,
    normalize_remote_availability,
    normalize_remote_download_status,
    normalize_remote_task_status,
    parse_result_evidence,
    parse_unambiguous_zulu_timestamp,
    protect_spreadsheet_text,
    sanitize_report_evidence,
    validate_report_agent_payload,
)
from wbcz.windows_agent import (
    AgentHttpResponse,
    AgentJob,
    AgentJobType,
    AgentResult,
    ProductionAgentTrueApiTransport,
)
from wbcz.windows_agent_runtime import (
    AgentCreateOutcomeUnresolved,
    DurableWindowsAgentExecutor,
    WindowsAgentReplayStore,
)
from wbcz_ui.live_true_api import TrueApiError
from wbcz_web.models import Base
from wbcz_web.models.reports import (
    ReportArtifactRecord,
    ReportArtifactUploadRecord,
    ReportJobEventRecord,
    ReportJobRecord,
    ReportSnapshotRecord,
)
from wbcz_web.repositories.reports import SqlReportRepository
from wbcz_web.services.reports import SynchronousReportExecutor


NOW = datetime(2026, 9, 18, 0, 0, tzinfo=timezone.utc)


class KeyProvider:
    def __init__(self, key: bytes = b"k" * 32) -> None:
        self.key = key
        self.calls: list[str] = []

    def get_key(self, key_version: str) -> bytes:
        self.calls.append(key_version)
        return self.key


class FakeTunnel:
    local_port = 16666

    def session_marker(self):
        return "marker"

    def assert_gost_session(self, marker):
        assert marker == "marker"


class FakeAudit:
    def __init__(self) -> None:
        self.records = []

    def record(self, **kwargs):
        self.records.append(kwargs)


class FakeResponse:
    def __init__(self, status: int, body: bytes, content_type: str = "application/json") -> None:
        self.status = status
        self._body = body
        self._offset = 0
        self.headers = {"Content-Type": content_type, "X-Request-ID": "req-1"}

    def read(self, size: int | None = None) -> bytes:
        if size is None:
            part = self._body[self._offset:]
            self._offset = len(self._body)
            return part
        part = self._body[self._offset:self._offset + size]
        self._offset += len(part)
        return part


class FakeConnection:
    def __init__(self, owner) -> None:
        self.owner = owner
        self.method = None
        self.target = None
        self.headers = {}

    def putrequest(self, method, target, skip_host=True):
        self.method = method
        self.target = target

    def putheader(self, name, value):
        self.headers[name] = value

    def endheaders(self, data=None):
        self.owner.calls.append((self.method, self.target, dict(self.headers), data))

    def getresponse(self):
        body = self.owner.response_body(self.method, self.target)
        content_type = "application/zip" if "/file" in (self.target or "") else "application/json"
        return FakeResponse(200, body, content_type)

    def close(self):
        pass


class FakeConnectionFactory:
    def __init__(self):
        self.calls = []

    def __call__(self, *args, **kwargs):
        return FakeConnection(self)

    @staticmethod
    def response_body(method, target):
        if "/file" in (target or ""):
            return b"PK\x03\x04REPORT-ZIP-CONTENT"
        if target and "/tasks/" in target and method == "GET":
            return b'{"id":"task-1","currentStatus":"PREPARATION"}'
        if target and "/results" in target:
            return b'{"list":[]}'
        if target and "/tasks" in target and method == "POST":
            return b'{"id":"task-1","currentStatus":"PREPARATION"}'
        return b'{"productGroups":[{"productGroupCode":1,"remainingCount":2}]}'


def create_payload(job_id: str = "rpt_123") -> dict:
    return validate_report_agent_payload(
        "REPORT_CREATE",
        {
            "local_report_job_id": job_id,
            "recipe": "FILTERED_CIS_REPORT",
            "product_group_code": "1",
            "filters": {
                "participant_inn": "1234567890",
                "package_type": ["UNIT"],
                "status": "INTRODUCED",
                "include_gtin": ["00000000000001"],
            },
        },
    )


def test_hard_security_gates_and_no_user_download() -> None:
    assert not PRODUCTION_TRUE_API_REPORTS_ENABLED
    assert not GENERIC_TRUE_API_REPORT_PROXY
    assert not ARBITRARY_REMOTE_URL_DOWNLOAD
    assert not CALLER_CONTROLLED_TRUE_API_HOST
    assert not CALLER_CONTROLLED_TRUE_API_PATH
    assert not CALLER_CONTROLLED_TRUE_API_METHOD
    assert not CALLER_CONTROLLED_FILESYSTEM_PATH
    assert not CALLER_CONTROLLED_STORAGE_KEY
    assert not PATH_TRAVERSAL_ALLOWED
    assert not REMOTE_RESULT_DELETE_ENABLED
    assert not USER_REPORT_DOWNLOAD_ENABLED
    assert FRONTEND_FREEZE_ACTIVE
    assert M12_DEPENDENCY_FOR_DOWNLOAD_AUTH == "PARTIAL"
    routes = Path("src/wbcz_web/api/routes.py").read_text(encoding="utf-8")
    assert '@router.get("/reports/' not in routes
    assert '@router.get("/report-artifacts/' not in routes


def test_local_report_catalog_has_all_required_types_and_blocker_labels() -> None:
    assert set(LOCAL_REPORT_CATALOG) == set(LocalReportType)
    assert len(LOCAL_REPORT_CATALOG) == 12
    m1 = LOCAL_REPORT_CATALOG[LocalReportType.CIS_INVENTORY_STORED_SNAPSHOT]
    assert m1.fixed_labels["observation_semantics"] == "LATEST_STORED_OBSERVATION"
    m7 = LOCAL_REPORT_CATALOG[LocalReportType.EDO_FOUNDATION_STATUS]
    assert m7.fixed_labels["blocker_reason"] == "M7_FULL_XML_WRITE_BLOCKED"
    m8 = LOCAL_REPORT_CATALOG[LocalReportType.SUZ_FOUNDATION_STATUS]
    assert m8.fixed_labels["blocker_reason"] == "M8_FULL_SUZ_WIRE_BLOCKED"
    assert m8.fixed_labels["km_decrypted"] is False
    m10 = LOCAL_REPORT_CATALOG[LocalReportType.OZON_FOUNDATION_STATUS]
    assert m10.fixed_labels["m10_wire_ready"] is False
    assert m10.fixed_labels["m10_executable_read_capabilities"] == "NONE"


def test_request_fingerprint_is_stable_and_only_accepted_fields_contribute() -> None:
    kwargs = dict(
        report_type="WB_RECONCILIATION_EVIDENCE",
        report_schema_version="1.0",
        participant_scope={"participant_inn": "1234567890"},
        normalized_filters={"state": "X", "nested": {"b": 2, "a": 1}},
        output_format="CSV",
        sensitivity_mode="BUSINESS_SENSITIVE",
        snapshot_policy_version="m11-v1",
    )
    a = canonical_request_fingerprint(**kwargs)
    b = canonical_request_fingerprint(**{
        **kwargs,
        "normalized_filters": {"nested": {"a": 1, "b": 2}, "state": "X"},
    })
    assert a == b and len(a) == 64
    assert canonical_request_fingerprint(**{**kwargs, "output_format": "JSON"}) != a


def test_snapshot_descriptor_hash_is_reproducible_and_source_sensitive() -> None:
    a = ReportSnapshotDescriptor(
        SnapshotStrategy.APPEND_ONLY_HIGH_WATER,
        NOW,
        ("M1",),
        ("agent_jobs",),
        {"max_job": "42"},
        {"state": "X"},
        "1.0",
        row_count=10,
    )
    b = ReportSnapshotDescriptor(
        SnapshotStrategy.APPEND_ONLY_HIGH_WATER,
        NOW,
        ("M1",),
        ("agent_jobs",),
        {"max_job": "42"},
        {"state": "X"},
        "1.0",
        row_count=10,
    )
    c = ReportSnapshotDescriptor(
        SnapshotStrategy.APPEND_ONLY_HIGH_WATER,
        NOW,
        ("M1",),
        ("agent_jobs",),
        {"max_job": "43"},
        {"state": "X"},
        "1.0",
        row_count=10,
    )
    assert a.descriptor_sha256 == b.descriptor_sha256
    assert a.descriptor_sha256 != c.descriptor_sha256


def test_reconciliation_keeps_remote_local_and_final_semantics_separate() -> None:
    row = build_reconciliation_row(
        participant_inn="1234567890",
        domain="WB_FBS",
        source="WB_API",
        source_fingerprint="a" * 64,
        local_operation_id="op-1",
        remote_id="remote-1",
        remote_raw_state="sold",
        local_project_state="DISTANCE_PENDING",
        final_result=LocalReconciliationResult.PENDING_EVIDENCE,
        evidence_missing=True,
        conflict_state=None,
        observed_at="2026-09-18T12:00:00",
    )
    assert row["remote_raw_state"] == "sold"
    assert row["local_project_state"] == "DISTANCE_PENDING"
    assert row["final_reconciliation_result"] == "PENDING_EVIDENCE"
    assert row["observed_at"] == "2026-09-18T12:00:00"


def test_timestamp_unknown_is_not_silently_assigned_timezone() -> None:
    assert format_timestamp_preserving_unknown(datetime(2026, 9, 18, 12, 0)) == "2026-09-18T12:00:00|TIMEZONE_UNKNOWN"
    assert format_timestamp_preserving_unknown("2026-09-18 12:00:00") == "2026-09-18 12:00:00"
    assert parse_unambiguous_zulu_timestamp("2026-09-18T12:00:00Z").utcoffset() == timedelta(0)
    assert parse_unambiguous_zulu_timestamp("2026-09-18T12:00:00") is None


def test_csv_formula_protection_unicode_quotes_newlines_and_no_truncation(tmp_path: Path) -> None:
    renderer = StreamingReportRenderer(temp_root=tmp_path, max_bytes=2_000_000)
    long_value = "Ж" * 40000
    rows = [{
        "id": "00000000000000001234",
        "text": 'Кириллица, "quote"\nnext',
        "formula": "=2+2",
        "at": "@cmd",
        "tab": "\tcmd",
        "cr": "\rcmd",
        "long": long_value,
    }]
    artifact = renderer.render(
        rows,
        columns=("id", "text", "formula", "at", "tab", "cr", "long"),
        output_format=ReportOutputFormat.CSV,
    )
    with artifact.path.open("r", encoding="utf-8", newline="") as handle:
        raw = handle.read()
    assert "00000000000000001234" in raw
    assert "Кириллица" in raw
    assert "'=2+2" in raw and "'@cmd" in raw and "'\tcmd" in raw and "'\rcmd" in raw
    assert long_value in raw
    artifact.path.unlink()


def test_xlsx_identifiers_are_text_and_oversize_cell_fails_not_truncates(tmp_path: Path) -> None:
    renderer = StreamingReportRenderer(temp_root=tmp_path, max_bytes=5_000_000)
    artifact = renderer.render(
        [{"id": "00000000000000001234", "uuid": "00000000-0000-0000-0000-000000000001", "formula": "+1"}],
        columns=("id", "uuid", "formula"),
        output_format=ReportOutputFormat.XLSX,
    )
    wb = load_workbook(artifact.path, read_only=True, data_only=False)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=False))
    assert rows[1][0].value == "00000000000000001234"
    assert rows[1][0].data_type == "s"
    assert rows[1][2].value == "'+1"
    wb.close()
    artifact.path.unlink()
    with pytest.raises(ReportFormatError):
        renderer.render(
            [{"cell": "x" * 32768}],
            columns=("cell",),
            output_format=ReportOutputFormat.XLSX,
        )


def test_chunked_sensitive_artifact_encryption_roundtrip_wrong_key_and_tamper(tmp_path: Path) -> None:
    source = tmp_path / "source.bin"
    source.write_bytes((b"CIS-CONTROL-\x1d-SECRET-" * 100000))
    store = FilesystemReportArtifactStore(
        tmp_path / "store",
        key_provider=KeyProvider(),
        key_version="test-v1",
        chunk_bytes=4096,
    )
    result = store.put_file(
        source,
        sensitive=True,
        aad_context={"artifact_id": "a1", "report_job_id": "r1"},
        safe_suffix=".enc",
    )
    assert result.encryption_version == ARTIFACT_ENCRYPTION_FORMAT
    stored = (tmp_path / "store" / result.storage_key).read_bytes()
    assert b"CIS-CONTROL" not in stored
    output = io.BytesIO()
    size, digest = store.decrypt_sensitive_to(
        result.storage_key, output,
        aad_context={"artifact_id": "a1", "report_job_id": "r1"},
    )
    assert output.getvalue() == source.read_bytes()
    assert size == source.stat().st_size
    assert digest == result.plaintext_sha256

    wrong = FilesystemReportArtifactStore(
        tmp_path / "store",
        key_provider=KeyProvider(b"x" * 32),
        key_version="test-v1",
        chunk_bytes=4096,
    )
    with pytest.raises(ReportSecurityError):
        wrong.decrypt_sensitive_to(
            result.storage_key, io.BytesIO(),
            aad_context={"artifact_id": "a1", "report_job_id": "r1"},
        )

    path = tmp_path / "store" / result.storage_key
    data = bytearray(path.read_bytes())
    data[-10] ^= 1
    path.write_bytes(data)
    with pytest.raises(ReportSecurityError):
        store.decrypt_sensitive_to(
            result.storage_key, io.BytesIO(),
            aad_context={"artifact_id": "a1", "report_job_id": "r1"},
        )


def test_filesystem_artifact_store_generates_key_and_rejects_path_traversal(tmp_path: Path) -> None:
    source = tmp_path / "x.txt"
    source.write_text("safe", encoding="utf-8")
    store = FilesystemReportArtifactStore(tmp_path / "store")
    result = store.put_file(source, sensitive=False, aad_context={}, safe_suffix=".txt")
    assert "/" not in result.storage_key and "\\" not in result.storage_key
    with pytest.raises(ReportSecurityError):
        store.inspect("../escape")


def test_artifact_reuse_requires_exact_snapshot_schema_format_and_ready() -> None:
    base = dict(
        old_request_fingerprint="a",
        new_request_fingerprint="a",
        old_snapshot_hash="s",
        new_snapshot_hash="s",
        old_schema_version="1",
        new_schema_version="1",
        old_format="CSV",
        new_format="CSV",
        ready=True,
        expired_or_deleted=False,
    )
    assert artifact_reuse_allowed(**base)
    assert not artifact_reuse_allowed(**{**base, "new_snapshot_hash": "other"})
    assert not artifact_reuse_allowed(**{**base, "ready": False})
    assert not artifact_reuse_allowed(**{**base, "expired_or_deleted": True})


def test_dispenser_registry_exact_capabilities_and_delete_disabled() -> None:
    expected = {
        DispenserCapabilityName.CREATE_EXPORT: ("POST", "/api/v3/true-api/dispenser/tasks", 15),
        DispenserCapabilityName.GET_TASK: ("GET", "/api/v3/true-api/dispenser/tasks/{taskId}", 5),
        DispenserCapabilityName.LIST_TASKS: ("GET", "/api/v3/true-api/dispenser/tasks", 5),
        DispenserCapabilityName.LIST_RESULTS: ("GET", "/api/v3/true-api/dispenser/results", 12),
        DispenserCapabilityName.DOWNLOAD_RESULT: ("GET", "/api/v3/true-api/dispenser/results/{resultId}/file", 12),
        DispenserCapabilityName.QUOTA_BY_TASK_TYPE: ("GET", "/api/v3/true-api/dispenser/tasktypes/available_count", 10),
        DispenserCapabilityName.QUOTA_BY_REPORT: ("GET", "/api/v3/true-api/dispenser/tasktypes/reports/{report_id}/available_count", 10),
    }
    for name, (method, path, rate) in expected.items():
        cap = DISPENSER_CAPABILITIES[name]
        assert cap.enabled and cap.method == method and cap.path_template == path and cap.requests_per_minute == rate
    delete = DISPENSER_CAPABILITIES[DispenserCapabilityName.REMOTE_RESULT_DELETE]
    assert delete.method == "DELETE" and not delete.enabled
    assert not REMOTE_RESULT_DELETE_ENABLED


def test_filtered_cis_recipe_is_only_enabled_create_recipe_and_has_no_raw_params_input() -> None:
    assert enabled_remote_recipes() == ("FILTERED_CIS_REPORT",)
    assert "DOCUMENTS_ERRORS" in disabled_remote_recipes()
    filters = FilteredCisFilters(
        "1234567890", ("UNIT", "LEVEL1"), "INTRODUCED",
        ("00000000000001",),
    )
    body = build_filtered_cis_create_body(product_group_code="1", filters=filters)
    assert body == {
        "format": "CSV",
        "name": "FILTERED_CIS_REPORT",
        "periodicity": "SINGLE",
        "productGroupCode": "1",
        "params": '{"includeGtin":["00000000000001"],"packageType":["UNIT","LEVEL1"],"participantInn":"1234567890","status":"INTRODUCED"}',
    }
    with pytest.raises(ReportContractError):
        validate_report_agent_payload(
            "REPORT_CREATE",
            {**create_payload(), "params": '{"raw":"forbidden"}'},
        )
    with pytest.raises(ReportContractError):
        validate_report_agent_payload(
            "REPORT_CREATE",
            {
                "local_report_job_id": "r1",
                "recipe": "DOCUMENTS_ERRORS",
                "product_group_code": "1",
                "filters": {},
            },
        )


def test_include_gtin_max_1000_and_status_required_in_restricted_recipe() -> None:
    FilteredCisFilters("1234567890", ("UNIT",), "INTRODUCED", tuple(str(i) for i in range(1000))).validate()
    with pytest.raises(ReportContractError):
        FilteredCisFilters("1234567890", ("UNIT",), "INTRODUCED", tuple(str(i) for i in range(1001))).validate()
    with pytest.raises(ReportContractError):
        FilteredCisFilters("1234567890", ("UNIT",), None, ()).validate()


def test_dispenser_specs_are_fixed_and_quota_report_is_typed() -> None:
    create = build_dispenser_spec("REPORT_CREATE", create_payload())
    assert create.method == "POST"
    assert create.target == "/api/v3/true-api/dispenser/tasks"
    assert create.json_body["name"] == "FILTERED_CIS_REPORT"

    task = build_dispenser_spec(
        "REPORT_TASK_GET",
        {"local_report_job_id": "r1", "task_id": "task-1", "product_group_code": "1"},
    )
    assert task.target == "/api/v3/true-api/dispenser/tasks/task-1?pg=1"

    results = build_dispenser_spec(
        "REPORT_RESULTS",
        {"local_report_job_id": "r1", "page": 0, "size": 100, "product_group_code": "1", "task_ids": ["task-1"]},
    )
    assert results.target.startswith("/api/v3/true-api/dispenser/results?")
    assert "task_ids=task-1" in results.target

    quota = build_dispenser_spec(
        "REPORT_QUOTA_ID",
        {"local_report_job_id": "r1", "report_id": "ppr-1955-1-gismt-c", "product_group_code": "1"},
    )
    assert quota.target == "/api/v3/true-api/dispenser/tasktypes/reports/ppr-1955-1-gismt-c/available_count?pg=1"

    with pytest.raises(ReportContractError):
        build_dispenser_spec(
            "REPORT_QUOTA_ID",
            {"local_report_job_id": "r1", "report_id": "../escape", "product_group_code": "1"},
        )


def test_remote_raw_statuses_preserve_known_and_unknown_exactly() -> None:
    for status in KNOWN_REMOTE_TASK_STATUSES:
        parsed = normalize_remote_task_status(status)
        assert parsed.raw == status and parsed.known
    future = normalize_remote_task_status("FUTURE_STATUS")
    assert future.raw == "FUTURE_STATUS" and not future.known

    for status in KNOWN_REMOTE_DOWNLOAD_STATUSES:
        assert normalize_remote_download_status(status).known
    assert not normalize_remote_download_status("FUTURE").known

    for status in KNOWN_REMOTE_AVAILABILITY:
        assert normalize_remote_availability(status).known
    assert not normalize_remote_availability("FUTURE").known


def test_result_parser_never_turns_remote_file_path_into_fetch_url() -> None:
    parsed = parse_result_evidence({
        "id": "res-1",
        "taskId": "task-1",
        "archiveSize": 123,
        "available": "AVAILABLE",
        "downloadStatus": "SUCCESS",
        "fileDeleteDate": "2026-09-20T12:00:00.000Z",
        "resultFileParts": [{"id": "part-1", "filePath": "../../secret"}],
    })
    assert parsed.result_id == "res-1"
    assert parsed.result_file_parts[0]["filePath"] == "../../secret"
    assert not hasattr(parsed, "fetch_url")


def transport(factory: FakeConnectionFactory, clock=lambda: 0.0) -> ProductionAgentTrueApiTransport:
    return ProductionAgentTrueApiTransport(
        tunnel=FakeTunnel(),
        audit=FakeAudit(),
        connection_factory=factory,
        report_rate_limiter=DispenserRateLimiter(),
        monotonic_clock=clock,
        production_true_api_reports=True,
    )


def test_actual_transport_enforces_create_15_per_minute_before_network() -> None:
    factory = FakeConnectionFactory()
    tx = transport(factory)
    payload = create_payload()
    for _ in range(REPORT_CREATE_RATE_PER_MINUTE):
        tx.m11_report("REPORT_CREATE", payload, bearer_token="BEARER-CANARY", rate_scope="1234567890")
    with pytest.raises(ReportRateLimitExceeded) as exc:
        tx.m11_report("REPORT_CREATE", payload, bearer_token="BEARER-CANARY", rate_scope="1234567890")
    assert exc.value.retry_after_seconds == pytest.approx(60.0)
    assert len(factory.calls) == 15
    assert all("BEARER-CANARY" in call[2]["Authorization"] for call in factory.calls)


def test_actual_transport_task_family_shared_5_per_minute() -> None:
    factory = FakeConnectionFactory()
    tx = transport(factory)
    get_payload = {"local_report_job_id": "r1", "task_id": "task-1", "product_group_code": "1"}
    list_payload = {"local_report_job_id": "r1", "page": 0, "size": 10, "product_group_code": "1"}
    for _ in range(4):
        tx.m11_report("REPORT_TASK_GET", get_payload, bearer_token="B", rate_scope="inn")
    tx.m11_report("REPORT_TASK_LIST", list_payload, bearer_token="B", rate_scope="inn")
    with pytest.raises(ReportRateLimitExceeded):
        tx.m11_report("REPORT_TASK_GET", get_payload, bearer_token="B", rate_scope="inn")
    assert len(factory.calls) == REPORT_TASK_RATE_PER_MINUTE


def test_actual_transport_results_and_download_share_12_per_minute() -> None:
    factory = FakeConnectionFactory()
    tx = transport(factory)
    payload = {"local_report_job_id": "r1", "page": 0, "size": 10, "product_group_code": "1", "task_ids": ["task-1"]}
    for _ in range(REPORT_RESULTS_RATE_PER_MINUTE):
        tx.m11_report("REPORT_RESULTS", payload, bearer_token="B", rate_scope="inn")
    with pytest.raises(ReportRateLimitExceeded):
        tx.m11_report_download(
            {
                "local_report_job_id": "r1",
                "result_id": "res-1",
                "result_part_id": None,
                "product_group_code": "1",
                "artifact_upload_id": "upl_" + "a" * 32,
                "expected_archive_size": None,
                "download_format": None,
            },
            bearer_token="B",
            rate_scope="inn",
            sink=io.BytesIO(),
        )
    assert len(factory.calls) == REPORT_RESULTS_RATE_PER_MINUTE


def test_actual_transport_quota_family_shared_10_per_minute_and_scopes_independent() -> None:
    factory = FakeConnectionFactory()
    tx = transport(factory)
    type_payload = {"local_report_job_id": "r1", "task_type_short_name": "FILTERED_CIS_REPORT", "product_group_code": "1"}
    report_payload = {"local_report_job_id": "r1", "report_id": "ppr-1955-1-gismt-c", "product_group_code": "1"}
    for _ in range(9):
        tx.m11_report("REPORT_QUOTA_TYPE", type_payload, bearer_token="B", rate_scope="inn-a")
    tx.m11_report("REPORT_QUOTA_ID", report_payload, bearer_token="B", rate_scope="inn-a")
    with pytest.raises(ReportRateLimitExceeded):
        tx.m11_report("REPORT_QUOTA_TYPE", type_payload, bearer_token="B", rate_scope="inn-a")
    tx.m11_report("REPORT_QUOTA_TYPE", type_payload, bearer_token="B", rate_scope="inn-b")
    assert len(factory.calls) == REPORT_QUOTA_RATE_PER_MINUTE + 1


def test_transport_default_production_report_gate_is_closed() -> None:
    factory = FakeConnectionFactory()
    tx = ProductionAgentTrueApiTransport(
        tunnel=FakeTunnel(),
        audit=FakeAudit(),
        connection_factory=factory,
    )
    with pytest.raises(Exception):
        tx.m11_report("REPORT_CREATE", create_payload(), bearer_token="B", rate_scope="inn")
    assert factory.calls == []


def test_recursive_security_sanitizer_removes_secret_and_full_marking_canaries() -> None:
    payload = {
        "Authorization": "Bearer SECRET-BEARER",
        "machine_token": "MACHINE-CANARY",
        "future": [{"key": "cis", "value": "FULL-CIS-CANARY"}],
        "nonSensitive": {"keep": "yes"},
    }
    safe = sanitize_report_evidence(payload)
    rendered = repr(safe)
    assert "SECRET-BEARER" not in rendered
    assert "MACHINE-CANARY" not in rendered
    assert "FULL-CIS-CANARY" not in rendered
    assert safe["nonSensitive"] == {"keep": "yes"}


def test_binary_artifact_ingress_streams_hashes_encrypts_and_replay_converges(tmp_path: Path) -> None:
    bindings = InMemoryUploadBindingStore()
    store = FilesystemReportArtifactStore(tmp_path / "store", key_provider=KeyProvider(), key_version="v1", chunk_bytes=16)
    ingress = BinaryArtifactIngress(binding_store=bindings, artifact_store=store, temp_root=tmp_path / "tmp", remote_download_byte_ceiling=10000)
    binding = ingress.prepare(
        report_job_id="rpt_1",
        remote_result_id="res-1",
        remote_result_part_id=None,
        product_group_code="1",
        expected_archive_size=12,
    )
    payload = b"PK0123456789"
    artifact = ingress.ingest(
        artifact_upload_id=binding.artifact_upload_id,
        report_job_id="rpt_1",
        remote_result_id="res-1",
        remote_result_part_id=None,
        chunks=[payload[:3], payload[3:]],
        observed_mime="application/zip",
    )
    assert artifact.byte_size == len(payload)
    assert artifact.sha256
    stored = (tmp_path / "store" / artifact.storage.storage_key).read_bytes()
    assert payload not in stored
    replay = ingress.ingest(
        artifact_upload_id=binding.artifact_upload_id,
        report_job_id="rpt_1",
        remote_result_id="res-1",
        remote_result_part_id=None,
        chunks=[payload],
        observed_mime="application/zip",
    )
    assert replay.artifact_id == artifact.artifact_id
    with pytest.raises(ArtifactIntegrityConflict):
        ingress.ingest(
            artifact_upload_id=binding.artifact_upload_id,
            report_job_id="rpt_1",
            remote_result_id="res-1",
            remote_result_part_id=None,
            chunks=[b"PKxxxxxxxxxx"],
            observed_mime="application/zip",
        )


def test_binary_artifact_ingress_binding_and_archive_size_fail_closed(tmp_path: Path) -> None:
    bindings = InMemoryUploadBindingStore()
    store = FilesystemReportArtifactStore(tmp_path / "store", key_provider=KeyProvider(), key_version="v1")
    ingress = BinaryArtifactIngress(binding_store=bindings, artifact_store=store, temp_root=tmp_path / "tmp")
    with pytest.raises(ReportSecurityError):
        ingress.ingest(
            artifact_upload_id="upl_" + "f" * 32,
            report_job_id="r",
            remote_result_id="res",
            remote_result_part_id=None,
            chunks=[b"x"],
            observed_mime=None,
        )
    binding = ingress.prepare(
        report_job_id="rpt_1",
        remote_result_id="res-1",
        remote_result_part_id="part-1",
        product_group_code="1",
        expected_archive_size=10,
    )
    with pytest.raises(ReportSecurityError):
        ingress.ingest(
            artifact_upload_id=binding.artifact_upload_id,
            report_job_id="WRONG",
            remote_result_id="res-1",
            remote_result_part_id="part-1",
            chunks=[b"1234567890"],
            observed_mime=None,
        )
    with pytest.raises(ArtifactIntegrityConflict):
        ingress.ingest(
            artifact_upload_id=binding.artifact_upload_id,
            report_job_id="rpt_1",
            remote_result_id="res-1",
            remote_result_part_id="part-1",
            chunks=[b"short"],
            observed_mime=None,
        )


class FakeSession:
    def bearer_token(self):
        return "BEARER-SHOULD-NOT-SERIALIZE"

    def observe_http_status(self, status):
        pass


class AmbiguousReportTransport:
    def __init__(self):
        self.calls = 0

    def m11_report(self, *args, **kwargs):
        self.calls += 1
        raise TrueApiError("network outcome unknown")


def report_create_job() -> AgentJob:
    return AgentJob(
        job_id="job_report_create_1",
        job_type=AgentJobType.REPORT_CREATE,
        operation_id="rpt_remote_1",
        pg="lp",
        expected_inn="1234567890",
        read_payload=create_payload("rpt_remote_1"),
    )


def test_ambiguous_create_is_durable_and_not_blindly_retried(tmp_path: Path) -> None:
    replay = WindowsAgentReplayStore(tmp_path / "replay.sqlite")
    fake = AmbiguousReportTransport()
    executor = DurableWindowsAgentExecutor(
        participant_inn="1234567890",
        transport=fake,
        session_manager=FakeSession(),
        document_signer=object(),
        replay_store=replay,
        production_write=False,
    )
    job = report_create_job()
    first = executor.execute(job)
    assert first.outcome == "REMOTE_CREATE_AMBIGUOUS"
    assert fake.calls == 1
    second = executor.execute(job)
    assert second.outcome == "REMOTE_CREATE_AMBIGUOUS"
    assert fake.calls == 1
    replay.close()


def test_unresolved_create_reservation_blocks_post_after_crash(tmp_path: Path) -> None:
    replay = WindowsAgentReplayStore(tmp_path / "replay.sqlite")
    fake = AmbiguousReportTransport()
    executor = DurableWindowsAgentExecutor(
        participant_inn="1234567890",
        transport=fake,
        session_manager=FakeSession(),
        document_signer=object(),
        replay_store=replay,
        production_write=False,
    )
    job = report_create_job()
    payload_sha = executor._report_create_payload_sha(job)
    replay.claim("m11-report-create:" + job.operation_id, payload_sha)
    with pytest.raises(AgentCreateOutcomeUnresolved):
        executor.execute(job)
    assert fake.calls == 0
    replay.close()


class FakeDownloadTransport:
    def m11_report_download(self, payload, *, bearer_token, rate_scope, sink):
        assert bearer_token == "BEARER-SHOULD-NOT-SERIALIZE"
        sink.write(b"ZIP-BINARY-SECRET")
        return {
            "http_status": 200,
            "byte_size": len(b"ZIP-BINARY-SECRET"),
            "sha256": __import__("hashlib").sha256(b"ZIP-BINARY-SECRET").hexdigest(),
            "content_type": "application/zip",
        }


def test_download_executor_uploads_file_and_result_json_contains_no_binary_base64(tmp_path: Path) -> None:
    from wbcz.windows_agent import WindowsAgentExecutor
    executor = WindowsAgentExecutor(
        participant_inn="1234567890",
        transport=FakeDownloadTransport(),
        session_manager=FakeSession(),
        document_signer=object(),
        production_write=False,
    )
    job = AgentJob(
        job_id="job_download",
        job_type=AgentJobType.REPORT_DOWNLOAD,
        operation_id="rpt_1",
        pg="lp",
        expected_inn="1234567890",
        read_payload=validate_report_agent_payload(
            "REPORT_DOWNLOAD",
            {
                "local_report_job_id": "rpt_1",
                "result_id": "res-1",
                "result_part_id": None,
                "product_group_code": "1",
                "artifact_upload_id": "upl_" + "a" * 32,
                "expected_archive_size": None,
                "download_format": None,
            },
        ),
    )
    uploaded = {}

    def upload(metadata, path):
        uploaded["metadata"] = dict(metadata)
        uploaded["bytes"] = path.read_bytes()

    result = executor.execute_report_download(job, upload=upload)
    assert uploaded["bytes"] == b"ZIP-BINARY-SECRET"
    rendered = json.dumps(result.safe_dict())
    assert "ZIP-BINARY-SECRET" not in rendered
    assert "WklQLUJJTkFSWS1TRUNSRVQ" not in rendered
    assert result.outcome == "REPORT_ARTIFACT_UPLOADED"


def _ensure_m11_report_tables(engine) -> None:
    names = (
        "report_snapshots", "report_jobs", "report_artifacts",
        "report_job_events", "report_artifact_uploads",
    )
    tables = [Base.metadata.tables[name] for name in names]
    Base.metadata.create_all(engine, tables=tables, checkfirst=True)


@pytest.mark.skipif(not os.getenv("WBCZ_TEST_DATABASE_URL"), reason="PostgreSQL integration database not configured")
def test_db_backed_report_claim_lease_recovery_bounded_attempts_and_no_double_worker() -> None:
    engine = create_engine(os.environ["WBCZ_TEST_DATABASE_URL"])
    _ensure_m11_report_tables(engine)
    with Session(engine) as db:
        repo = SqlReportRepository(db)
        row = repo.create_job(
            origin="LOCAL",
            participant_inn="1234567890",
            report_type=LocalReportType.OZON_FOUNDATION_STATUS.value,
            report_schema_version="1.0",
            output_format="CSV",
            sensitivity_class="NORMAL",
            filters_sanitized={},
            sensitive_filter_ref=None,
            request_fingerprint_sha256="a" * 64,
            requested_by_user_id=None,
        )
        repo.queue(row.id)
        base = datetime.now(timezone.utc)
        first = repo.claim_next(worker_id="w1", lease_seconds=30, max_attempts=2, now=base)
        assert first and first.attempt_count == 1
        assert repo.claim_next(worker_id="w2", lease_seconds=30, max_attempts=2, now=base + timedelta(seconds=10)) is None
        second = repo.claim_next(worker_id="w2", lease_seconds=30, max_attempts=2, now=base + timedelta(seconds=31))
        assert second and second.attempt_count == 2 and second.lease_owner == "w2"
        assert repo.claim_next(worker_id="w3", lease_seconds=30, max_attempts=2, now=base + timedelta(seconds=62)) is None
        assert db.get(ReportJobRecord, row.id).state == "FAILED"
        db.rollback()
    engine.dispose()


@pytest.mark.skipif(not os.getenv("WBCZ_TEST_DATABASE_URL"), reason="PostgreSQL integration database not configured")
def test_synchronous_local_worker_generates_ready_artifact_without_business_mutation(tmp_path: Path) -> None:
    engine = create_engine(os.environ["WBCZ_TEST_DATABASE_URL"])
    _ensure_m11_report_tables(engine)
    with Session(engine) as db:
        repo = SqlReportRepository(db)
        row = repo.create_job(
            origin="LOCAL",
            participant_inn="1234567890",
            report_type=LocalReportType.OZON_FOUNDATION_STATUS.value,
            report_schema_version="1.0",
            output_format="CSV",
            sensitivity_class="NORMAL",
            filters_sanitized={},
            sensitive_filter_ref=None,
            request_fingerprint_sha256="b" * 64,
            requested_by_user_id=None,
        )
        repo.queue(row.id)
        claimed = repo.claim_next(worker_id="sync", lease_seconds=60, max_attempts=3)
        assert claimed
        store = FilesystemReportArtifactStore(tmp_path / "store", key_provider=KeyProvider(), key_version="v1")
        executor = SynchronousReportExecutor(
            db,
            artifact_store=store,
            temp_root=tmp_path / "tmp",
            participant_inn="1234567890",
        )
        result = executor.execute_claimed(claimed)
        assert result.row_count == 1
        assert db.get(ReportJobRecord, row.id).state == "READY"
        artifact = db.get(ReportArtifactRecord, result.artifact_id)
        assert artifact and artifact.state == "READY" and artifact.artifact_role == "OUTPUT"
        snapshot = db.get(ReportSnapshotRecord, db.get(ReportJobRecord, row.id).snapshot_id)
        assert snapshot and snapshot.descriptor_sha256 == result.snapshot_hash
        db.rollback()
    engine.dispose()


def test_persistence_models_have_no_bearer_token_pin_private_key_or_plaintext_cis_columns() -> None:
    for model in (ReportJobRecord, ReportSnapshotRecord, ReportArtifactRecord, ReportJobEventRecord, ReportArtifactUploadRecord):
        columns = {x.lower() for x in model.__table__.columns.keys()}
        assert not {
            "bearer", "bearer_token", "uuidtoken", "token", "machine_token", "pin",
            "private_key", "certificate_private_key", "cis", "kiz", "sgtin", "full_cis",
        } & columns
    assert {"storage_key", "sha256", "byte_size", "encryption_metadata"} <= set(ReportArtifactRecord.__table__.columns.keys())


def test_migration_head_contract_and_agent_types_are_bounded() -> None:
    migration = Path("migrations/versions/0012_m11_reports.py").read_text(encoding="utf-8")
    assert 'revision = "0012_m11_reports"' in migration
    assert 'down_revision = "0011_m10_ozon"' in migration
    for job_type in M11_AGENT_JOB_TYPES:
        assert job_type in migration
    for forbidden in ("TRUE_API_HTTP", "REPORT_HTTP", "ARBITRARY_REQUEST"):
        assert forbidden not in migration


def test_no_generic_true_api_proxy_arbitrary_url_or_user_download_surface() -> None:
    source = Path("src/wbcz/m11_reports.py").read_text(encoding="utf-8")
    windows = Path("src/wbcz/windows_agent.py").read_text(encoding="utf-8")
    routes = Path("src/wbcz_web/api/agent_routes.py").read_text(encoding="utf-8")
    for forbidden in (
        "GenericTrueApiReportProxy", "execute(method, url", "remote_url",
        "caller_storage_key", "caller_path",
    ):
        assert forbidden not in source
    assert "report-artifacts/{artifact_upload_id}" in routes
    assert "request.stream()" in routes
    assert "request.body()" not in routes[routes.find("agent_report_artifact_ingress"):]
    assert "base64.b64encode" not in windows[windows.find("def execute_report_download"):windows.find("def _cis_check")]
