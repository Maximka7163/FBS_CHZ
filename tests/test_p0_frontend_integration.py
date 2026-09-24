from __future__ import annotations

from collections import Counter
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import json
import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from wbcz.document_assembler import OrganisationType
from wbcz.models import Decision, Event, Operation
from wbcz.windows_agent import AgentResult
from wbcz.write_pipeline import ExactDocumentBuilder, InvalidWriteOperation
from wbcz_web.auth import hash_password
from wbcz_web.config import WebConfig
from wbcz_web.main import create_app
from wbcz_web.models import (
    AgentJobRecord,
    Base,
    CheckRecord,
    ControlRun,
    ImportRecord,
    ImportRow,
    User,
    WriteOperationRecord,
)
from wbcz_web.repositories import ImportRepository
from wbcz_web.services.agent_orchestration import AgentControlService, CONTROL_CIS, CONTROL_CIS_BATCH, WRITE
from wbcz_web.services.document_orchestration import AgentOrchestrationBroker
from wbcz_web.services.imports import FileImportService, event_to_record
from wbcz_web.services.authorization import BootstrapService
from wbcz_web.services.tenant import bind_tenant_scope
from wbcz_web.services.workspace import (
    BulkActionUnavailable,
    FbsDryRunWriteBlocked,
    bulk_preview,
    execute_bulk_actions,
    workspace_overview,
)


DB_URL = os.getenv("WBCZ_TEST_DATABASE_URL")
OWN = "1234567890"
TOKEN = "frontend-integration-machine-token-0123456789"
TEST_AUDIT_KEY = "synthetic-test-audit-pseudonym-key-000000000001"
TEST_AUDIT_KEY_ID = "frontend-test-audit-v1"
FIAS = "11111111-2222-3333-4444-555555555555"
KPP = "123456789"
CIS_PREFIX = "010290089707781021"
ROOT = Path(__file__).parents[1]


def valid_cis(label: str) -> str:
    value = CIS_PREFIX + label
    assert 18 <= len(value) <= 74
    return value


def event(label: str, operation: Operation = Operation.SALE) -> Event:
    return Event(
        kiz=valid_cis(label),
        task_number=f"task-{label}",
        sticker=f"sticker-{label}",
        operation=operation,
        occurred_at=datetime(2026, 8, 19, 12, 0, tzinfo=timezone.utc),
        receipt_number=f"receipt-{label}",
        fiscal_drive_number="7380440903834317",
        amount=Decimal("1901.75"),
        currency="RUB",
        legal_entity_sale=False,
    )


def config(
    database_url: str,
    *,
    agent: bool = True,
    org: bool = True,
    dry_run: bool = False,
    real_read: bool = False,
) -> WebConfig:
    kwargs = dict(
        database_url=database_url,
        own_inn=OWN,
        environment="test",
        agent_enabled=agent,
        agent_machine_token=TOKEN if agent else None,
        agent_job_lease_seconds=30,
        agent_poll_initial_seconds=1,
        agent_poll_max_seconds=4,
        agent_poll_max_attempts=3,
        true_api_write_enabled=False,
        true_api_real_read_enabled=real_read,
        fbs_dry_run_only=dry_run,
        audit_pseudonym_key=TEST_AUDIT_KEY,
        audit_pseudonym_key_id=TEST_AUDIT_KEY_ID,
    )
    if org:
        kwargs.update(
            organisation_type=OrganisationType.LEGAL_ENTITY,
            activity_fias_id=FIAS,
            activity_kpp=KPP,
            remote_sale_return_paid=False,
        )
    return WebConfig(**kwargs).validate_for_startup()


@pytest.fixture
def pg_factory():
    if not DB_URL:
        pytest.skip("WBCZ_TEST_DATABASE_URL requires PostgreSQL")
    engine = create_engine(DB_URL, future=True)
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    try:
        yield factory
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def seed_import(db, events: list[Event], *, password: str | None = None):
    user = User(
        username="operator",
        password_hash=hash_password(password) if password else "not-used",
        is_active=True,
        is_admin=True,
    )
    db.add(user)
    db.flush()
    imported = ImportRecord(
        fingerprint=("f" * 63) + "1",
        filename="wb-fbs.xlsx",
        imported_by=user.id,
        new_events=len(events),
        duplicate_events=0,
        rejected_rows=0,
        row_count=len(events),
        unique_kiz=len({item.kiz for item in events}),
        sales=sum(item.operation is Operation.SALE for item in events),
        returns=sum(item.operation is Operation.RETURN for item in events),
        dated=len(events),
        undated=0,
    )
    db.add(imported)
    db.flush()
    for index, item in enumerate(events, start=2):
        db.add(event_to_record(item))
        db.flush()
        db.add(ImportRow(import_id=imported.id, row_number=index, event_id=item.event_id))
    db.flush()
    return user, imported


def add_check(db, imported: ImportRecord, user: User, item: Event, decision: Decision, reason: str, *, error: str | None = None):
    run = ControlRun(import_id=imported.id, user_id=user.id, mode="AUTO", provider="test")
    db.add(run)
    db.flush()
    snapshot = None
    if decision is Decision.READY_TO_WITHDRAW:
        snapshot = {"status": "IN_CIRCULATION", "statusEx": None, "withdrawReason": None, "ownerInn": OWN, "productGroup": "lp"}
    elif decision is Decision.READY_TO_RETURN:
        snapshot = {"status": "WITHDRAWN", "statusEx": None, "withdrawReason": "DISTANCE", "ownerInn": OWN, "productGroup": "lp"}
    db.add(
        CheckRecord(
            run_id=run.id,
            event_id=item.event_id,
            source="test",
            snapshot=snapshot,
            decision=decision.value,
            reason=reason,
            error=error,
        )
    )
    db.flush()


def workbook_bytes() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "КИЗ"
    sheet.append([
        "№ задания",
        "Стикер",
        "КИЗ",
        "Номер чека",
        "Стоимость",
        "Валюта",
        "Номер фискального накопителя",
        "Дата",
        "Тип операции",
        "Признак продажи юрлицу",
    ])
    sheet.append([
        "100",
        "200",
        valid_cis("UPLOAD"),
        "300",
        1901.75,
        "RUB",
        "7380440903834317",
        "12:34:56 19.08.2026",
        "Продажа",
        "нет",
    ])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def login(client: TestClient, password: str) -> None:
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/api/auth/login",
        headers={"X-CSRF-Token": csrf},
        json={"username": "operator", "password": password},
    )
    assert response.status_code == 200


def test_frontend_browser_boundary_has_no_business_persistence_or_direct_true_api():
    sources = "\n".join(
        (ROOT / "frontend" / "src" / name).read_text(encoding="utf-8")
        for name in ("main.ts", "api.ts", "workflow.ts")
    )
    assert "localStorage" not in sources
    assert "sessionStorage" not in sources
    assert "markirovka.crpt.ru" not in sources
    assert "/api/agent/" not in sources
    assert "WBCZ_AGENT_MACHINE_TOKEN" not in sources
    assert "CryptoPro" not in sources
    assert '"/api/' in sources
    assert "READY_TO_WITHDRAW" not in (ROOT / "frontend" / "src" / "workflow.ts").read_text(encoding="utf-8")
    assert "READY_TO_RETURN" not in (ROOT / "frontend" / "src" / "workflow.ts").read_text(encoding="utf-8")


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_workspace_groups_states_and_bulk_is_backend_authoritative(pg_factory):
    rows = [
        event("WITHDRAW"),
        event("RETURN", Operation.RETURN),
        event("MANUAL"),
        event("ERROR"),
        event("DONE"),
        event("NOACTION"),
    ]
    with pg_factory() as db:
        user, imported = seed_import(db, rows)
        decisions = [
            (Decision.READY_TO_WITHDRAW, "SALE_IN_CIRCULATION"),
            (Decision.READY_TO_RETURN, "RETURN_WITHDRAWN_DISTANCE"),
            (Decision.MANUAL_REVIEW, "OWNER_MISMATCH"),
            (Decision.ERROR, "STATE_LOOKUP_OR_NORMALIZATION_FAILED"),
            (Decision.ALREADY_DONE, "SALE_ALREADY_WITHDRAWN_DISTANCE"),
            (Decision.NO_ACTION, "TEST_NO_ACTION"),
        ]
        for item, (decision, reason) in zip(rows, decisions):
            add_check(db, imported, user, item, decision, reason, error="TrueApiError" if decision is Decision.ERROR else None)
        view = workspace_overview(db, config(DB_URL), imported.id)
        assert view["filters"] == {"ALL": 6, "READY": 2, "PROCESSING": 0, "ATTENTION": 1, "DONE": 2, "ERROR": 1}
        assert view["bulk"]["eligible_count"] == 2
        assert view["bulk"]["withdraw_count"] == 1
        assert view["bulk"]["return_count"] == 1
        by_decision = {item["decision"]: item for item in view["items"]}
        assert by_decision[Decision.READY_TO_WITHDRAW.value]["action_label"] == "Вывести из оборота"
        assert by_decision[Decision.READY_TO_RETURN.value]["action_label"] == "Возврат в оборот"
        assert by_decision[Decision.MANUAL_REVIEW.value]["ready_for_bulk"] is False
        assert by_decision[Decision.ERROR.value]["ready_for_bulk"] is False
        assert by_decision[Decision.ALREADY_DONE.value]["ready_for_bulk"] is False
        assert by_decision[Decision.NO_ACTION.value]["ready_for_bulk"] is False
        preview = bulk_preview(db, imported.id)
        assert preview["eligible_count"] == 2
        assert preview["excluded_count"] == 4


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_true_api_write_disabled_is_hard_stop_for_bulk_and_prepare(pg_factory):
    item = event("WRITE-GATE")
    cfg = config(DB_URL)
    with pg_factory() as db:
        user, imported = seed_import(db, [item])
        add_check(db, imported, user, item, Decision.READY_TO_WITHDRAW, "SALE_IN_CIRCULATION")
        with pytest.raises(BulkActionUnavailable, match="TRUE_API_WRITE_DISABLED"):
            execute_bulk_actions(db, cfg, imported.id, user.id)
        document = ExactDocumentBuilder.from_json_value({"write_gate": False})
        with pytest.raises(InvalidWriteOperation, match="TRUE_API_WRITE_DISABLED"):
            AgentOrchestrationBroker(db, cfg).prepare_approved_write(
                item.event_id,
                document,
                decision=Decision.READY_TO_WITHDRAW,
            )
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_agent_control_result_is_decision_only_while_write_gate_off(pg_factory):
    item = event("CONTROL")
    cfg = config(DB_URL)
    with pg_factory() as db:
        user, imported = seed_import(db, [item])
        response = AgentControlService(db, cfg).run(imported.id, user.id, "AUTO", [item.event_id])
        assert response["pending"] == 1
        job = db.scalar(select(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS))
        assert job is not None
        broker = AgentOrchestrationBroker(db, cfg)
        broker.submit_result(
            TOKEN,
            AgentResult(
                job.job_id,
                job.operation_id,
                "CIS_CHECKED",
                cises=({
                    "cis": item.kiz,
                    "status": "IN_CIRCULATION",
                    "statusEx": None,
                    "withdrawReason": None,
                    "ownerInn": OWN,
                    "productGroup": "lp",
                },),
            ),
        )
        check = ImportRepository(db).latest_check(item.event_id)
        assert check is not None and check.decision == Decision.READY_TO_WITHDRAW.value
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0
        with pytest.raises(BulkActionUnavailable, match="TRUE_API_WRITE_DISABLED"):
            execute_bulk_actions(db, cfg, imported.id, user.id)
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
@pytest.mark.parametrize(
    ("operations", "latest_state", "latest_decision"),
    [
        ((Operation.SALE, Operation.RETURN), "WITHDRAWN", Decision.READY_TO_RETURN),
        ((Operation.RETURN, Operation.SALE), "IN_CIRCULATION", Decision.READY_TO_WITHDRAW),
    ],
)
def test_sequence_control_queues_only_latest_event(
    pg_factory, operations, latest_state, latest_decision,
):
    base_event = event("SEQ", operations[0])
    older = replace(base_event, operation=operations[0], occurred_at=datetime(2026, 8, 19, 10, 0, tzinfo=timezone.utc))
    latest = replace(base_event, operation=operations[1], task_number="task-latest", occurred_at=datetime(2026, 8, 19, 11, 0, tzinfo=timezone.utc))
    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        user, imported = seed_import(db, [older, latest])
        response = AgentControlService(db, cfg).run(
            imported.id, user.id, "AUTO", [older.event_id, latest.event_id]
        )
        assert response["pending"] == 1
        assert response["checked"] == 1
        jobs = list(db.scalars(select(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS)))
        assert len(jobs) == 1
        assert jobs[0].event_id == latest.event_id
        old_check = ImportRepository(db).latest_check(older.event_id)
        assert old_check.decision == Decision.NO_ACTION.value
        assert old_check.reason == "SUPERSEDED_BY_LATER_WB_EVENT"

        result_item = {
            "cis": latest.kiz,
            "status": latest_state,
            "statusEx": None,
            "withdrawReason": "DISTANCE" if latest_state == "WITHDRAWN" else None,
            "ownerInn": OWN,
            "productGroup": "lp",
        }
        AgentOrchestrationBroker(db, cfg).submit_result(
            TOKEN,
            AgentResult(jobs[0].job_id, jobs[0].operation_id, "CIS_CHECKED", cises=(result_item,)),
        )
        latest_check = ImportRepository(db).latest_check(latest.event_id)
        assert latest_check.decision == latest_decision.value
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_three_event_sequence_marks_all_older_events_superseded(pg_factory):
    base_event = event("SEQ3", Operation.SALE)
    events = [
        replace(base_event, operation=Operation.SALE, task_number="one", occurred_at=datetime(2026, 8, 19, 9, 0, tzinfo=timezone.utc)),
        replace(base_event, operation=Operation.RETURN, task_number="two", occurred_at=datetime(2026, 8, 19, 10, 0, tzinfo=timezone.utc)),
        replace(base_event, operation=Operation.SALE, task_number="three", occurred_at=datetime(2026, 8, 19, 11, 0, tzinfo=timezone.utc)),
    ]
    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        user, imported = seed_import(db, events)
        response = AgentControlService(db, cfg).run(
            imported.id, user.id, "AUTO", [item.event_id for item in events]
        )
        assert response["pending"] == 1
        assert response["counts"][Decision.NO_ACTION.value] == 2
        for old in events[:2]:
            check = ImportRepository(db).latest_check(old.event_id)
            assert check.decision == Decision.NO_ACTION.value
            assert check.reason == "SUPERSEDED_BY_LATER_WB_EVENT"
        job = db.scalar(select(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS))
        assert job.event_id == events[-1].event_id


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
@pytest.mark.parametrize("ambiguous_kind", ["missing", "equal"])
def test_ambiguous_event_order_never_queues_true_api_check(pg_factory, ambiguous_kind):
    first = event("AMB", Operation.SALE)
    if ambiguous_kind == "missing":
        second = replace(first, operation=Operation.RETURN, task_number="second", occurred_at=None)
    else:
        second = replace(first, operation=Operation.RETURN, task_number="second")
    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        user, imported = seed_import(db, [first, second])
        response = AgentControlService(db, cfg).run(
            imported.id, user.id, "AUTO", [first.event_id, second.event_id]
        )
        assert response["pending"] == 0
        assert response["counts"] == {Decision.MANUAL_REVIEW.value: 2}
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS)) == 0
        for item in (first, second):
            check = ImportRepository(db).latest_check(item.event_id)
            assert check.decision == Decision.MANUAL_REVIEW.value
            assert check.reason == "HISTORY_ORDER_AMBIGUOUS"


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_dry_run_blocks_bulk_and_prepare_write_without_creating_state(pg_factory):
    item = event("DRYGUARD")
    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        user, imported = seed_import(db, [item])
        add_check(db, imported, user, item, Decision.READY_TO_WITHDRAW, "SALE_IN_CIRCULATION")
        with pytest.raises(FbsDryRunWriteBlocked, match="FBS_DRY_RUN_ONLY"):
            execute_bulk_actions(db, cfg, imported.id, user.id)
        document = ExactDocumentBuilder.from_json_value({"dry_run": True})
        with pytest.raises(InvalidWriteOperation, match="FBS_DRY_RUN_ONLY"):
            AgentOrchestrationBroker(db, cfg).prepare_approved_write(
                item.event_id,
                document,
                decision=Decision.READY_TO_WITHDRAW,
            )
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_workspace_evidence_is_explicit_whitelist_and_dryrun_runtime(pg_factory):
    item = event("EVIDENCE")
    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        user, imported = seed_import(db, [item])
        run = ControlRun(import_id=imported.id, user_id=user.id, mode="AUTO", provider="test")
        db.add(run); db.flush()
        db.add(CheckRecord(
            run_id=run.id,
            event_id=item.event_id,
            source="windows-agent-true-api",
            snapshot={
                "status": "IN_CIRCULATION",
                "statusEx": None,
                "withdrawReason": None,
                "ownerInn": OWN,
                "productGroup": "lp",
                "fetched_at": "2026-09-22T06:00:00+00:00",
                "raw_internal": "MUST_NOT_LEAK",
                "token": "MUST_NOT_LEAK",
            },
            decision=Decision.READY_TO_WITHDRAW.value,
            reason="SALE_IN_CIRCULATION",
        ))
        db.flush()
        view = workspace_overview(db, cfg, imported.id)
        row = view["items"][0]
        assert {
            key: row[key]
            for key in ("status", "statusEx", "withdrawReason", "ownerInn", "owner_match", "productGroup", "source", "reason_code")
        } == {
            "status": "IN_CIRCULATION",
            "statusEx": None,
            "withdrawReason": None,
            "ownerInn": OWN,
            "owner_match": None,
            "productGroup": "lp",
            "source": "windows-agent-true-api",
            "reason_code": "SALE_IN_CIRCULATION",
        }
        assert row["fetched_at"] == "2026-09-22T06:00:00+00:00"
        assert row["checked_at"] is not None
        serialized = json.dumps(view)
        assert "MUST_NOT_LEAK" not in serialized
        assert "raw_internal" not in serialized
        assert '"token"' not in serialized
        assert view["runtime"]["fbs_dry_run_only"] is True
        assert view["runtime"]["production_write_enabled"] is False
        assert view["runtime"]["true_api_real_read_enabled"] is False
        assert view["runtime"]["control_provider"] == "mock-dry-run"


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_http_dry_run_bulk_bypass_is_typed_409_and_creates_no_write(pg_factory):
    password = "very-secure-dryrun-password"
    cfg = config(DB_URL, agent=True, org=True, dry_run=True)
    with pg_factory() as db:
        _,org,participant,_=BootstrapService(db).bootstrap(
            username="operator",password=password,organisation_name="Dry Run",
            participant_inn=OWN,
        )
        org_id,participant_id=org.id,participant.id
        db.commit()
    app = create_app(cfg, session_factory=pg_factory)
    with TestClient(app) as client:
        login(client, password)
        csrf = client.get("/api/auth/csrf").json()["csrf_token"]
        client.post("/api/security/scope", headers={"X-CSRF-Token":csrf}, json={"organisation_id":org_id,"participant_id":participant_id})
        uploaded = client.post("/api/files", headers={"X-CSRF-Token":csrf}, files={"file":("wb.xlsx",workbook_bytes(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        import_id = uploaded.json()["id"]
        response = client.post(f"/api/files/{import_id}/bulk-actions", headers={"X-CSRF-Token":csrf}, json={"confirm":True})
        assert response.status_code == 409
        assert response.json()["detail"]["code"] == "FBS_DRY_RUN_ONLY"
    with pg_factory() as db:
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_http_mock_control_ignores_agent_presence_when_live_read_is_disabled(pg_factory):
    password = "very-secure-mock-provider-password"
    cfg = config(DB_URL, agent=True, org=True, dry_run=True, real_read=False)
    with pg_factory() as db:
        _,org,participant,_=BootstrapService(db).bootstrap(
            username="operator",password=password,organisation_name="Mock Provider",
            participant_inn=OWN,
        )
        org_id,participant_id=org.id,participant.id
        db.commit()
    app = create_app(cfg, session_factory=pg_factory)
    with TestClient(app) as client:
        login(client, password)
        csrf = client.get("/api/auth/csrf").json()["csrf_token"]
        client.post("/api/security/scope", headers={"X-CSRF-Token":csrf}, json={"organisation_id":org_id,"participant_id":participant_id})
        uploaded = client.post("/api/files", headers={"X-CSRF-Token":csrf}, files={"file":("wb.xlsx",workbook_bytes(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        import_id = uploaded.json()["id"]
        response = client.post(
            f"/api/files/{import_id}/control",
            headers={"X-CSRF-Token":csrf},
            json={"mode":"AUTO","event_ids":None},
        )
        assert response.status_code == 200
        assert response.json()["provider"] == "mock"
        assert response.json()["checked"] == 1
        capabilities = client.get("/api/capabilities").json()
        assert capabilities["true_api"] == "offline-dry-run"
        assert capabilities["windows_bridge"] is True
    with pg_factory() as db:
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_http_live_read_only_control_uses_agent_without_enabling_writes(pg_factory):
    password = "very-secure-live-provider-password"
    cfg = config(DB_URL, agent=True, org=True, dry_run=True, real_read=True)
    with pg_factory() as db:
        _,org,participant,_=BootstrapService(db).bootstrap(
            username="operator",password=password,organisation_name="Live Provider",
            participant_inn=OWN,
        )
        org_id,participant_id=org.id,participant.id
        db.commit()
    app = create_app(cfg, session_factory=pg_factory)
    with TestClient(app) as client:
        login(client, password)
        csrf = client.get("/api/auth/csrf").json()["csrf_token"]
        client.post("/api/security/scope", headers={"X-CSRF-Token":csrf}, json={"organisation_id":org_id,"participant_id":participant_id})
        uploaded = client.post("/api/files", headers={"X-CSRF-Token":csrf}, files={"file":("wb.xlsx",workbook_bytes(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        import_id = uploaded.json()["id"]
        response = client.post(
            f"/api/files/{import_id}/control",
            headers={"X-CSRF-Token":csrf},
            json={"mode":"AUTO","event_ids":None},
        )
        assert response.status_code == 200
        assert response.json()["provider"] == "windows-agent"
        assert response.json()["pending"] == 1
        view = client.get(f"/api/files/{import_id}/workspace").json()
        assert view["runtime"]["control_provider"] == "live-read-only"
        assert view["runtime"]["production_write_enabled"] is False
        capabilities = client.get("/api/capabilities").json()
        assert capabilities["true_api"] == "windows-agent"
        assert capabilities["windows_bridge"] is True
    with pg_factory() as db:
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS_BATCH)) == 1
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_exact_duplicate_rows_create_one_event_and_one_cis_check(pg_factory):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "КИЗ"
    headers = (
        "№ задания", "Стикер", "КИЗ", "Номер чека", "Стоимость", "Валюта",
        "Номер фискального накопителя", "Дата", "Тип операции", "Признак продажи юрлицу",
    )
    sheet.append(list(headers))
    row = [
        "dup-task", "dup-sticker", valid_cis("DUP"), "dup-check", 100, "RUB",
        "7380440903834317", "12:00:00 19.08.2026", "Продажа", "нет",
    ]
    sheet.append(row)
    sheet.append(row)
    stream = BytesIO()
    wb.save(stream)
    wb.close()

    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        user, org, participant, _ = BootstrapService(db).bootstrap(
            username="operator-dup",
            password="very-secure-duplicate-password",
            organisation_name="Duplicate Regression",
            participant_inn=OWN,
        )
        bind_tenant_scope(
            db,
            organisation_id=org.id,
            participant_id=participant.id,
            user_id=user.id,
            role="ADMIN",
        )
        imported = FileImportService(db).import_xlsx("duplicate.xlsx", stream.getvalue(), user.id)
        assert imported.new_events == 1
        assert imported.duplicate_events == 1
        event_ids = [row.event_id for row in ImportRepository(db).ordered_event_records(imported.id)]
        response = AgentControlService(db, cfg).run(imported.id, user.id, "AUTO", event_ids)
        assert response["pending"] == 1
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS)) == 1


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_live_ready_is_blocked_after_newer_overlapping_wb_event_before_any_write_state(pg_factory):
    sale = event("STALE-READY", Operation.SALE)
    read_cfg = config(DB_URL, real_read=True)
    with pg_factory() as db:
        user, imported = seed_import(db, [sale])
        response = AgentControlService(db, read_cfg).run(
            imported.id, user.id, "AUTO", [sale.event_id]
        )
        assert response["pending"] == 1
        job = db.scalar(select(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS))
        assert job is not None

        AgentOrchestrationBroker(db, read_cfg).submit_result(
            TOKEN,
            AgentResult(
                job.job_id,
                job.operation_id,
                "CIS_CHECKED",
                cises=({
                    "cis": sale.kiz,
                    "status": "IN_CIRCULATION",
                    "statusEx": None,
                    "withdrawReason": None,
                    "ownerInn": OWN,
                    "productGroup": "lp",
                },),
            ),
        )
        check = ImportRepository(db).latest_check(sale.event_id)
        assert check is not None
        assert check.source == "windows-agent-true-api"
        assert check.decision == Decision.READY_TO_WITHDRAW.value

        newer_return = replace(
            sale,
            operation=Operation.RETURN,
            task_number="task-newer-return",
            sticker="sticker-newer-return",
            occurred_at=sale.occurred_at + timedelta(hours=1),
            receipt_number=None,
            fiscal_drive_number=None,
        )
        newer_import = ImportRecord(
            fingerprint=("f" * 63) + "2",
            filename="wb-fbs-newer.xlsx",
            imported_by=user.id,
            new_events=1,
            duplicate_events=0,
            rejected_rows=0,
            row_count=1,
            unique_kiz=1,
            sales=0,
            returns=1,
            dated=1,
            undated=0,
        )
        db.add(newer_import)
        db.flush()
        db.add(event_to_record(newer_return))
        db.flush()
        db.add(ImportRow(import_id=newer_import.id, row_number=2, event_id=newer_return.event_id))
        db.flush()

        write_cfg = replace(
            read_cfg,
            fbs_dry_run_only=False,
            true_api_write_enabled=True,
            true_api_real_read_enabled=True,
        )

        with pytest.raises(BulkActionUnavailable, match="SUPERSEDED_BY_LATER_WB_EVENT"):
            execute_bulk_actions(db, write_cfg, imported.id, user.id)
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0

        with pytest.raises(InvalidWriteOperation, match="SUPERSEDED_BY_LATER_WB_EVENT"):
            AgentOrchestrationBroker(db, write_cfg).prepare_approved_write(
                sale.event_id,
                ExactDocumentBuilder.from_json_value({"stale_ready": True}),
                decision=Decision.READY_TO_WITHDRAW,
            )
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_result_time_sequence_recheck_prevents_stale_ready_after_newer_import(pg_factory):
    older = event("RACE", Operation.SALE)
    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        user, imported = seed_import(db, [older])
        response = AgentControlService(db, cfg).run(imported.id, user.id, "AUTO", [older.event_id])
        assert response["pending"] == 1
        job = db.scalar(select(AgentJobRecord).where(AgentJobRecord.purpose == CONTROL_CIS))
        assert job is not None

        newer = replace(
            older,
            operation=Operation.RETURN,
            task_number="race-later",
            occurred_at=older.occurred_at + timedelta(minutes=1),
        )
        db.add(event_to_record(newer))
        db.flush()

        AgentOrchestrationBroker(db, cfg).submit_result(
            TOKEN,
            AgentResult(
                job.job_id,
                job.operation_id,
                "CIS_CHECKED",
                cises=({
                    "cis": older.kiz,
                    "status": "IN_CIRCULATION",
                    "statusEx": None,
                    "withdrawReason": None,
                    "ownerInn": OWN,
                    "productGroup": "lp",
                },),
            ),
        )
        check = ImportRepository(db).latest_check(older.event_id)
        assert check.decision == Decision.NO_ACTION.value
        assert check.reason == "SUPERSEDED_BY_LATER_WB_EVENT"
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_238_row_upload_agent_results_and_workspace_regression(pg_factory, wb_regression_rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "КИЗ"
    headers = (
        "№ задания", "Стикер", "КИЗ", "Номер чека", "Стоимость", "Валюта",
        "Номер фискального накопителя", "Дата", "Тип операции", "Признак продажи юрлицу",
    )
    sheet.append(list(headers))
    for row in wb_regression_rows:
        sheet.append([row.get(name) for name in headers])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    cfg = config(DB_URL, dry_run=True)
    with pg_factory() as db:
        db.info["audit_pseudonym_key"] = TEST_AUDIT_KEY.encode("utf-8")
        db.info["audit_pseudonym_key_id"] = TEST_AUDIT_KEY_ID
        user, org, participant, _ = BootstrapService(db).bootstrap(
            username="operator238",
            password="very-secure-238-password",
            organisation_name="Regression 238",
            participant_inn=OWN,
        )
        bind_tenant_scope(
            db,
            organisation_id=org.id,
            participant_id=participant.id,
            user_id=user.id,
            role="ADMIN",
        )
        imported = FileImportService(db).import_xlsx("wb-regression-238.xlsx", stream.getvalue(), user.id)
        assert imported.row_count == 238
        assert imported.new_events == 238
        assert imported.duplicate_events == 0

        event_ids = [row.event_id for row in ImportRepository(db).ordered_event_records(imported.id)]
        response = AgentControlService(db, cfg).run(imported.id, user.id, "AUTO", event_ids)
        assert response["pending"] == 238
        assert response["checked"] == 0
        jobs = list(db.scalars(
            select(AgentJobRecord)
            .where(AgentJobRecord.purpose == CONTROL_CIS)
            .order_by(AgentJobRecord.created_at, AgentJobRecord.job_id)
        ))
        assert len(jobs) == 238

        broker = AgentOrchestrationBroker(db, cfg)
        imports = ImportRepository(db)
        for job in jobs:
            row = imports.event(job.event_id)
            assert row is not None
            event_value = Event.from_dict(dict(row.payload))
            withdrawn = event_value.operation is Operation.RETURN
            broker.submit_result(
                TOKEN,
                AgentResult(
                    job.job_id,
                    job.operation_id,
                    "CIS_CHECKED",
                    cises=({
                        "cis": event_value.kiz,
                        "status": "WITHDRAWN" if withdrawn else "IN_CIRCULATION",
                        "statusEx": None,
                        "withdrawReason": "DISTANCE" if withdrawn else None,
                        "ownerInn": OWN,
                        "productGroup": "lp",
                    },),
                ),
            )

        view = workspace_overview(db, cfg, imported.id)
        assert len(view["items"]) == 238
        decisions = Counter(item["decision"] for item in view["items"])
        assert decisions == {
            Decision.READY_TO_WITHDRAW.value: 63,
            Decision.READY_TO_RETURN.value: 162,
            Decision.MANUAL_REVIEW.value: 13,
        }
        assert all(item["source"] == "windows-agent-true-api" for item in view["items"])
        assert all(item["owner_match"] is True for item in view["items"])
        assert all(item["productGroup"] == "lp" for item in view["items"])
        assert db.scalar(select(func.count()).select_from(WriteOperationRecord)) == 0
        assert db.scalar(select(func.count()).select_from(AgentJobRecord).where(AgentJobRecord.purpose == WRITE)) == 0


@pytest.mark.skipif(not DB_URL, reason="PostgreSQL required")
def test_http_upload_restore_history_and_machine_secret_non_exposure(pg_factory):
    password = "very-secure-frontend-password"
    cfg = config(DB_URL, agent=True, org=True)
    with pg_factory() as db:
        _,org,participant,_=BootstrapService(db).bootstrap(
            username="operator",password=password,organisation_name="P0 Frontend Regression",
            participant_inn=OWN,
        )
        org_id,participant_id=org.id,participant.id
        db.commit()
    app = create_app(cfg, session_factory=pg_factory)
    with TestClient(app) as client:
        login(client, password)
        csrf = client.get("/api/auth/csrf").json()["csrf_token"]
        switched=client.post(
            "/api/security/scope",
            headers={"X-CSRF-Token":csrf},
            json={"organisation_id":org_id,"participant_id":participant_id},
        )
        assert switched.status_code==200
        uploaded = client.post(
            "/api/files",
            headers={"X-CSRF-Token": csrf},
            files={"file": ("wb-upload.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert uploaded.status_code == 200
        import_id = uploaded.json()["id"]
        home = client.get("/api/workspace")
        assert home.status_code == 200
        assert home.json()["active_import_id"] == import_id
        assert home.json()["history"][0]["id"] == import_id
        view = client.get(f"/api/files/{import_id}/workspace")
        assert view.status_code == 200
        body = view.json()
        assert body["file"]["filename"] == "wb-upload.xlsx"
        assert body["runtime"]["production_write_enabled"] is False
        combined = json.dumps({"home": home.json(), "view": body}, ensure_ascii=False)
        assert TOKEN not in combined
        assert "agent_machine_token" not in combined
        assert "bearer" not in combined.lower()
        bad_confirm = client.post(
            f"/api/files/{import_id}/bulk-actions",
            headers={"X-CSRF-Token": csrf},
            json={"confirm": False},
        )
        assert bad_confirm.status_code == 422


def test_production_write_default_remains_off():
    config_value = WebConfig(database_url="sqlite:///ignored", own_inn=OWN)
    assert config_value.true_api_write_enabled is False
    assert config_value.fbs_dry_run_only is False


def test_fbs_dry_run_env_flag_and_frontend_execution_gate(monkeypatch):
    monkeypatch.setenv("WBCZ_ENV", "test")
    monkeypatch.setenv("WBCZ_FBS_DRY_RUN_ONLY", "true")
    monkeypatch.setenv(
        "WBCZ_DATABASE_URL",
        "postgresql+psycopg://postgres:postgres@127.0.0.1:5432/wbcz_test",
    )
    monkeypatch.setenv("WBCZ_OWN_INN", OWN)
    monkeypatch.delenv("WBCZ_AGENT_ENABLED", raising=False)
    monkeypatch.delenv("WBCZ_PRINTING_ENABLED", raising=False)
    monkeypatch.delenv("WBCZ_PRINT_EXECUTION_ENABLED", raising=False)
    cfg = WebConfig.from_env()
    assert cfg.fbs_dry_run_only is True
    assert cfg.true_api_write_enabled is False
    source = (ROOT / "frontend" / "src" / "main.ts").read_text(encoding="utf-8")
    assert "view.runtime.fbs_dry_run_only" in source
    assert "DRY RUN — реальные действия отключены" in source
